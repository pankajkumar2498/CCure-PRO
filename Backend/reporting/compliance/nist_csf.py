"""
reporting/compliance/nist_csf.py
─────────────────────────────────────────────────────────────────────────────
The NIST CSF 2.0 Organizational Profile, as PDF and DOCX.

Moved unchanged from ``routers/compliance.py`` (G6, C8 of
docs/architecture/MODULE_INDEPENDENCE_ANALYSIS.md): a renderer takes the
controls and their rolled-up verdicts the router read, and returns the
document; it reads no table and imports no compliance code.
"""

import io
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


# ─────────────────────────────────────────────────────────────────────────────
# NIST CSF 2.0 Organizational Profile
#
# The CSF report used the generic Statement of Applicability generator, which
# is an ISO artefact: five columns headed "Control ID / Name / Category / Scan
# Mode / Verdict", with `Category` printing the raw Function code. Nothing in
# it named a Function, nothing grouped by Category, and there was no per-
# Function score — so the one thing every CSF reader looks for, which Function
# is weakest, could not be read off it. NIST calls this artefact a Current
# Profile, and it is organised Function ▸ Category ▸ Subcategory.
# ─────────────────────────────────────────────────────────────────────────────

# Verdict -> (label, hex colour). `not_assessed` and `manual_review` are kept
# apart on purpose: "nobody has looked at this" and "a human still has to sign
# it off" are different states, and collapsing them is what lets an unscanned
# framework read as a failing one.
_CSF_VERDICT_DISPLAY = {
    "pass": ("Compliant", "#16A34A"),
    "exception": ("Compliant (exception)", "#0891B2"),
    "fail": ("Non-compliant", "#DC2626"),
    "pending_reverify": ("Non-compliant (re-verify)", "#EA580C"),
    "remediation_applied": ("Partially compliant", "#CA8A04"),
    "manual_review": ("Awaiting review", "#7C3AED"),
    "not_applicable": ("Not applicable", "#64748B"),
    "not_assessed": ("Not assessed", "#94A3B8"),
}


def _csf_verdict(findings_map: dict, control_id: str) -> str:
    value = findings_map.get(control_id)
    if value is None:
        return "not_assessed"
    if isinstance(value, str):
        return value
    return str(getattr(value, "status", "not_assessed"))


def _generate_nist_csf_pdf(
    controls: list, findings_map: dict, meta: dict, grouped: list
) -> io.BytesIO:
    """NIST CSF 2.0 Current Profile, grouped Function ▸ Category ▸ Subcategory."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="NIST CSF 2.0 Current Profile",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CSFTitle",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0A2E5C"),
    )
    sub_style = ParagraphStyle(
        "CSFSub",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=9.5,
        textColor=colors.HexColor("#475569"),
    )
    fn_style = ParagraphStyle(
        "CSFFunction",
        parent=styles["Heading2"],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#0A2E5C"),
    )
    cat_style = ParagraphStyle(
        "CSFCategory",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9.5,
        textColor=colors.HexColor("#1E293B"),
    )
    cell = ParagraphStyle("CSFCell", parent=styles["Normal"], fontSize=7.5, leading=9.5)
    hdr = ParagraphStyle(
        "CSFHdr",
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white,
    )
    kpi = ParagraphStyle(
        "CSFKpi",
        fontSize=8,
        leading=11,
        alignment=1,
        textColor=colors.HexColor("#334155"),
    )

    story: list = [
        Paragraph("NIST Cybersecurity Framework 2.0 — Current Profile", title_style),
        Paragraph(
            f"{meta.get('organization', 'Organisation')} · "
            f"Generated {meta.get('generated_at', '')}"
            + (f" · Scope: {meta['scope']}" if meta.get("scope") else ""),
            sub_style,
        ),
        Spacer(1, 4 * mm),
        Paragraph(
            "Scores are computed over applicable Subcategories only. A Subcategory "
            "marked Not applicable is excluded from its Function's denominator; one "
            "that is Not assessed or Awaiting review remains in the denominator and "
            "counts against the score, because coverage is part of the posture.",
            sub_style,
        ),
        Spacer(1, 5 * mm),
    ]

    # Function summary — the headline a CSF reader is looking for.
    summary_rows = [
        [
            Paragraph("Function", hdr),
            Paragraph("Categories", hdr),
            Paragraph("Subcategories", hdr),
            Paragraph("Applicable", hdr),
            Paragraph("Assessed", hdr),
            Paragraph("Compliant", hdr),
            Paragraph("Score", hdr),
        ]
    ]
    for code, name, (pct, passed, applicable, assessed), categories in grouped:
        subcount = sum(len(rows) for _, _, rows in categories)
        summary_rows.append(
            [
                Paragraph(f"<b>{code}</b> {name}", cell),
                Paragraph(str(len(categories)), cell),
                Paragraph(str(subcount), cell),
                Paragraph(str(applicable), cell),
                Paragraph(str(assessed), cell),
                Paragraph(str(passed), cell),
                Paragraph(f"<b>{pct}%</b>", cell),
            ]
        )
    summary = Table(
        summary_rows,
        colWidths=[95 * mm, 22 * mm, 26 * mm, 22 * mm, 21 * mm, 22 * mm, 20 * mm],
        repeatRows=1,
    )
    summary.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A2E5C")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F8FAFC")],
                ),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(summary)

    total = len(controls)
    all_verdicts = [_csf_verdict(findings_map, c.id) for c in controls]
    applicable_all = [v for v in all_verdicts if v != "not_applicable"]
    passed_all = sum(1 for v in applicable_all if v in ("pass", "exception"))
    overall = (
        round(passed_all / len(applicable_all) * 100, 1) if applicable_all else 0.0
    )
    story.append(Spacer(1, 4 * mm))
    story.append(
        Table(
            [
                [
                    Paragraph(
                        f"SUBCATEGORIES<br/><font size=13><b>{total}</b></font>", kpi
                    ),
                    Paragraph(
                        "COMPLIANT<br/>"
                        f"<font size=13 color='#16A34A'><b>{passed_all}</b></font>",
                        kpi,
                    ),
                    Paragraph(
                        "NON-COMPLIANT<br/><font size=13 color='#DC2626'><b>"
                        f"{sum(1 for v in all_verdicts if v in ('fail', 'pending_reverify'))}"
                        "</b></font>",
                        kpi,
                    ),
                    Paragraph(
                        "AWAITING REVIEW<br/><font size=13 color='#7C3AED'><b>"
                        f"{sum(1 for v in all_verdicts if v == 'manual_review')}"
                        "</b></font>",
                        kpi,
                    ),
                    Paragraph(
                        "NOT ASSESSED<br/><font size=13 color='#94A3B8'><b>"
                        f"{sum(1 for v in all_verdicts if v == 'not_assessed')}"
                        "</b></font>",
                        kpi,
                    ),
                    Paragraph(
                        "NOT APPLICABLE<br/><font size=13 color='#64748B'><b>"
                        f"{sum(1 for v in all_verdicts if v == 'not_applicable')}"
                        "</b></font>",
                        kpi,
                    ),
                    Paragraph(
                        f"OVERALL<br/><font size=13 color='#0A2E5C'><b>{overall}%</b></font>",
                        kpi,
                    ),
                ]
            ],
            colWidths=[38 * mm] * 6 + [38 * mm],
            style=TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                    ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#CBD5E1")),
                    ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E2E8F0")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            ),
        )
    )

    for code, name, (pct, passed, applicable, assessed), categories in grouped:
        story.append(Spacer(1, 7 * mm))
        story.append(
            Paragraph(
                f"{code} — {name} &nbsp;·&nbsp; {pct}% "
                f"({passed}/{applicable} applicable Subcategories compliant)",
                fn_style,
            )
        )
        for cat_code, cat_name, rows in categories:
            story.append(Spacer(1, 2 * mm))
            story.append(Paragraph(f"{cat_code} · {cat_name}", cat_style))
            data = [
                [
                    Paragraph("Subcategory", hdr),
                    Paragraph("Outcome", hdr),
                    Paragraph("Assessment", hdr),
                    Paragraph("Status", hdr),
                    Paragraph("Owner / Reviewed", hdr),
                ]
            ]
            for c in rows:
                verdict = _csf_verdict(findings_map, c.id)
                label, colour = _CSF_VERDICT_DISPLAY.get(
                    verdict, (verdict.replace("_", " ").title(), "#334155")
                )
                rolled = findings_map.get(c.id)
                owner = getattr(rolled, "owner", None) or ""
                reviewed = getattr(rolled, "reviewed_at", None)
                data.append(
                    [
                        Paragraph(f"<b>{c.control_number}</b>", cell),
                        Paragraph(c.title or "", cell),
                        Paragraph(
                            (
                                "Automated"
                                if (c.scannability or "manual") == "auto"
                                else "Attestation"
                            ),
                            cell,
                        ),
                        Paragraph(
                            f"<font color='{colour}'><b>{label}</b></font>", cell
                        ),
                        Paragraph(
                            f"{owner}<br/>{reviewed.strftime('%d %b %Y') if reviewed else ''}",
                            cell,
                        ),
                    ]
                )
            tbl = Table(
                data,
                colWidths=[22 * mm, 128 * mm, 24 * mm, 40 * mm, 34 * mm],
                repeatRows=1,
            )
            tbl.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
                        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        (
                            "ROWBACKGROUNDS",
                            (0, 1),
                            (-1, -1),
                            [colors.white, colors.HexColor("#F8FAFC")],
                        ),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            story.append(tbl)

    doc.build(story)
    buffer.seek(0)
    return buffer


def _generate_nist_csf_docx(
    controls: list, findings_map: dict, meta: dict, grouped: list
) -> io.BytesIO:
    """The same Current Profile as a Word document."""
    try:
        from docx import Document as _Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt, RGBColor
    except ImportError:
        return io.BytesIO(b"python-docx is not installed.")

    document = _Document()
    heading = document.add_paragraph()
    run = heading.add_run("NIST Cybersecurity Framework 2.0 — Current Profile")
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(10, 46, 92)

    meta_p = document.add_paragraph()
    meta_run = meta_p.add_run(
        f"{meta.get('organization', 'Organisation')} · "
        f"Generated {meta.get('generated_at', '')}"
        + (f" · Scope: {meta['scope']}" if meta.get("scope") else "")
    )
    meta_run.font.italic = True
    meta_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    document.add_paragraph("Function summary").runs[0].bold = True
    summary = document.add_table(rows=1, cols=5)
    summary.style = "Table Grid"
    for i, name in enumerate(
        ["Function", "Subcategories", "Applicable", "Compliant", "Score"]
    ):
        summary.rows[0].cells[i].text = name
        for p in summary.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
    for code, name, (pct, passed, applicable, _assessed), categories in grouped:
        cells = summary.add_row().cells
        cells[0].text = f"{code} — {name}"
        cells[1].text = str(sum(len(rows) for _, _, rows in categories))
        cells[2].text = str(applicable)
        cells[3].text = str(passed)
        cells[4].text = f"{pct}%"

    for code, name, (pct, passed, applicable, _assessed), categories in grouped:
        document.add_paragraph()
        fn_p = document.add_paragraph()
        fn_run = fn_p.add_run(
            f"{code} — {name} · {pct}% ({passed}/{applicable} applicable compliant)"
        )
        fn_run.font.size = Pt(13)
        fn_run.font.bold = True
        fn_run.font.color.rgb = RGBColor(10, 46, 92)

        for cat_code, cat_name, rows in categories:
            cat_p = document.add_paragraph()
            cat_run = cat_p.add_run(f"{cat_code} · {cat_name}")
            cat_run.font.bold = True

            table = document.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            for i, name_ in enumerate(
                ["Subcategory", "Outcome", "Assessment", "Status"]
            ):
                table.rows[0].cells[i].text = name_
                for p in table.rows[0].cells[i].paragraphs:
                    for r in p.runs:
                        r.font.bold = True
            for c in rows:
                verdict = _csf_verdict(findings_map, c.id)
                label = _CSF_VERDICT_DISPLAY.get(
                    verdict, (verdict.replace("_", " ").title(), "")
                )[0]
                cells = table.add_row().cells
                cells[0].text = str(c.control_number or "")
                cells[1].text = str(c.title or "")
                cells[2].text = (
                    "Automated"
                    if (c.scannability or "manual") == "auto"
                    else "Attestation"
                )
                cells[3].text = label

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
