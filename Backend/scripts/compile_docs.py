"""
Backend/scripts/compile_docs.py
───────────────────────────────────────────────────────────────────────────────
Modular Markdown to Word (.docx) Documentation Compiler for CCure Audit PRO.

Classes:
  - MarkdownParser: Parses Markdown files into structured document tokens.
  - DocxRenderer: Renders structured tokens into an enterprise Word document with
    visual callout boxes, dynamic table column widths, inline formatting,
    hyperlinks, heading orphan control (keep_with_next), row split prevention (cantSplit),
    and header row repeating (tblHeader).
"""

import os
import re
import sys

try:
    import docx
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    print("python-docx is not installed. Installing python-docx...")
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor


class MarkdownParser:
    """Parses markdown text files into token dictionaries."""

    def __init__(self, file_path):
        self.file_path = file_path

    def parse(self):
        if not os.path.isfile(self.file_path):
            print(f"Warning: File {self.file_path} not found. Skipping.")
            return []

        with open(self.file_path, "r", encoding="utf-8", errors="ignore") as f:
            lines = f.readlines()

        tokens = []
        in_code_block = False
        code_text = []
        in_table = False
        table_rows = []

        for line in lines:
            raw_line = line.rstrip("\r\n")
            stripped = raw_line.strip()

            # Code Block Toggle
            if stripped.startswith("```"):
                if in_code_block:
                    tokens.append({"type": "code_block", "text": "".join(code_text)})
                    in_code_block = False
                    code_text = []
                else:
                    if in_table and table_rows:
                        tokens.append({"type": "table", "rows": table_rows})
                        in_table = False
                        table_rows = []
                    in_code_block = True
                continue

            if in_code_block:
                code_text.append(raw_line + "\n")
                continue

            # Table Detection
            if stripped.startswith("|") and stripped.endswith("|"):
                cells = [c.strip() for c in stripped.split("|")[1:-1]]
                if all(re.match(r"^:?-+:?$", c) for c in cells if c):
                    continue  # Skip delimiter row
                in_table = True
                table_rows.append(cells)
                continue
            else:
                if in_table and table_rows:
                    tokens.append({"type": "table", "rows": table_rows})
                    in_table = False
                    table_rows = []

            # Headings
            match_heading = re.match(r"^(#{1,6})\s+(.*)", stripped)
            if match_heading:
                level = len(match_heading.group(1))
                text = match_heading.group(2)
                tokens.append({"type": "heading", "level": level, "text": text})
                continue

            # Bullet List (Preserve indentation)
            match_bullet = re.match(r"^(\s*)([-*+])\s+(.*)", raw_line)
            if match_bullet:
                indent = len(match_bullet.group(1))
                level = indent // 2
                tokens.append(
                    {
                        "type": "list_bullet",
                        "level": level,
                        "text": match_bullet.group(3),
                    }
                )
                continue

            # Numbered List
            match_num = re.match(r"^(\s*)(\d+)\.\s+(.*)", raw_line)
            if match_num:
                indent = len(match_num.group(1))
                level = indent // 2
                tokens.append(
                    {"type": "list_number", "level": level, "text": match_num.group(3)}
                )
                continue

            # Alert / Callout Block
            if stripped.startswith(">"):
                alert_text = stripped[1:].strip()
                tokens.append({"type": "alert", "text": alert_text})
                continue

            if stripped == "":
                continue

            # Paragraph
            tokens.append({"type": "paragraph", "text": stripped})

        # EOF Flushing Safeguard
        if in_code_block and code_text:
            tokens.append({"type": "code_block", "text": "".join(code_text)})
        if in_table and table_rows:
            tokens.append({"type": "table", "rows": table_rows})

        return tokens


class DocxRenderer:
    """Renders structured tokens into an enterprise Word (.docx) document."""

    def __init__(self, output_path):
        self.output_path = output_path
        self.doc = Document()
        self._configure_document_styles()

    def _configure_document_styles(self):
        for section in self.doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        self._add_header_footer()

    def _add_header_footer(self):
        section = self.doc.sections[0]
        header_p = section.header.paragraphs[0]
        header_p.text = "CCure-Audit-PRO — Technical Documentation Suite"
        header_p.alignment = 2  # Right-aligned
        header_p.style.font.size = Pt(8.5)
        header_p.style.font.italic = True

        footer_p = section.footer.paragraphs[0]
        footer_p.text = "CCure-Audit-PRO (Production Build 2026) | Confidential"
        footer_p.alignment = 1  # Center-aligned
        footer_p.style.font.size = Pt(8.5)

    def render_cover_page(self):
        p = self.doc.add_paragraph()
        p.alignment = 1
        run_title = p.add_run("CCure-Audit-PRO\n")
        run_title.font.size = Pt(28)
        run_title.font.bold = True
        run_title.font.name = "Arial"

        run_sub = p.add_run("Enterprise Technical Documentation Suite\n\n")
        run_sub.font.size = Pt(18)
        run_sub.font.italic = True
        run_sub.font.name = "Arial"

        meta_p = self.doc.add_paragraph()
        meta_p.alignment = 1
        meta_run = meta_p.add_run(
            "Compiled Directly from Authoritative Source Codebase\nDate: July 2026 | Version 2.4.0-Prod"
        )
        meta_run.font.size = Pt(11)
        meta_run.font.name = "Arial"

        # Table of Contents Notice
        toc_p = self.doc.add_paragraph()
        toc_p.paragraph_format.space_before = Pt(36)
        toc_p.alignment = 1
        toc_run = toc_p.add_run(
            "[ Document Table of Contents — Refresh with F9 upon opening ]"
        )
        toc_run.font.size = Pt(9.5)
        toc_run.font.italic = True
        toc_run.font.color.rgb = RGBColor(128, 128, 128)

        self.doc.add_page_break()

    def add_section_header(self, title):
        p = self.doc.add_paragraph()
        run = p.add_run(title)
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102)
        p.style = self.doc.styles["Heading 1"]
        p.paragraph_format.keep_with_next = True

    def render_tokens(self, tokens):
        for token in tokens:
            ttype = token["type"]

            if ttype == "heading":
                self._render_heading(token["level"], token["text"])
            elif ttype == "paragraph":
                self._render_paragraph(token["text"])
            elif ttype == "code_block":
                self._render_code_block(token["text"])
            elif ttype == "table":
                self._render_table(token["rows"])
            elif ttype == "list_bullet":
                self._render_list(
                    token["text"], is_numbered=False, level=token.get("level", 0)
                )
            elif ttype == "list_number":
                self._render_list(
                    token["text"], is_numbered=True, level=token.get("level", 0)
                )
            elif ttype == "alert":
                self._render_alert(token["text"])

    def _render_heading(self, level, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.keep_with_next = True  # Prevent orphaned headings

        # Shift level by +1 (Section title is Heading 1)
        target_level = level + 1
        if target_level == 2:
            p.style = self.doc.styles["Heading 2"]
            run = p.add_run(text)
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)
        elif target_level == 3:
            p.style = self.doc.styles["Heading 3"]
            run = p.add_run(text)
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)
        else:
            p.style = self.doc.styles["Heading 4"]
            run = p.add_run(text)
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 51, 51)

    def _render_paragraph(self, text):
        p = self.doc.add_paragraph()
        self._add_formatted_text(p, text)

    def _render_code_block(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)

    def _render_list(self, text, is_numbered=False, level=0):
        style_name = "List Number" if is_numbered else "List Bullet"
        p = self.doc.add_paragraph(style=style_name)
        p.paragraph_format.left_indent = Inches(0.25 + (0.25 * level))
        self._add_formatted_text(p, text)

    def _render_alert(self, text):
        alert_type = "NOTE"
        match = re.match(r"^\[!(NOTE|TIP|IMPORTANT|WARNING|CAUTION)\]\s*(.*)", text)
        if match:
            alert_type = match.group(1).upper()
            text = match.group(2)

        styles = {
            "NOTE": {"color": "0066CC", "fill": "F0F8FF"},
            "TIP": {"color": "008000", "fill": "F4FBF4"},
            "IMPORTANT": {"color": "008000", "fill": "F4FBF4"},
            "WARNING": {"color": "CC3300", "fill": "FFF5EE"},
            "CAUTION": {"color": "CC3300", "fill": "FFF5EE"},
        }
        cfg = styles.get(alert_type, styles["NOTE"])

        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = 1
        cell = table.cell(0, 0)
        self._set_cell_background(cell, cfg["fill"])

        # Set 4.5pt left border
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = OxmlElement("w:tcBorders")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "36")
        left.set(qn("w:color"), cfg["color"])
        tcBorders.append(left)

        for b_name in ["top", "bottom", "right"]:
            b = OxmlElement(f"w:{b_name}")
            b.set(qn("w:val"), "none")
            tcBorders.append(b)
        tcPr.append(tcBorders)

        p = cell.paragraphs[0]
        p.paragraph_format.left_indent = Inches(0.1)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)

        tag_run = p.add_run(f"{alert_type}: ")
        tag_run.font.bold = True
        tag_run.font.name = "Arial"
        tag_run.font.size = Pt(9.5)
        tag_run.font.color.rgb = RGBColor.from_string(cfg["color"])

        self._add_formatted_text(p, text)
        self.doc.add_paragraph()  # spacing

    def _render_table(self, rows):
        if not rows:
            return

        num_cols = max(len(r) for r in rows)
        word_table = self.doc.add_table(rows=len(rows), cols=num_cols)
        word_table.style = "Table Grid"

        # Calculate dynamic column widths
        col_max_lens = [0] * num_cols
        for row in rows:
            for c_idx, cell_value in enumerate(row):
                if c_idx < num_cols:
                    col_max_lens[c_idx] = max(col_max_lens[c_idx], len(cell_value))

        total_len = sum(col_max_lens)
        col_widths = []
        if total_len > 0:
            for l in col_max_lens:
                w = (l / total_len) * 6.5
                w = max(0.6, min(w, 4.0))  # Clamp between 0.6" and 4.0"
                col_widths.append(Inches(w))
        else:
            col_widths = [Inches(6.5 / num_cols)] * num_cols

        # Header Repeat (tblHeader)
        header_trPr = word_table.rows[0]._tr.get_or_add_trPr()
        header_trPr.append(OxmlElement("w:tblHeader"))

        for r_idx, row_cells in enumerate(rows):
            row_tr = word_table.rows[r_idx]
            # Prevent row split across pages (cantSplit)
            trPr = row_tr._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))

            for c_idx, cell_value in enumerate(row_cells):
                if c_idx >= len(row_tr.cells):
                    continue
                cell = row_tr.cells[c_idx]
                cell.width = col_widths[c_idx]
                cell.text = ""
                p = cell.paragraphs[0]

                if r_idx == 0:
                    self._set_cell_background(cell, "003366")
                    self._add_formatted_text(p, cell_value, is_header=True)
                else:
                    if r_idx % 2 == 0:
                        self._set_cell_background(cell, "F2F2F2")
                    self._add_formatted_text(p, cell_value, is_header=False)

        self.doc.add_paragraph()  # Spacing after table

    def _add_formatted_text(self, paragraph, text, is_header=False):
        tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))", text)
        for token in tokens:
            if not token:
                continue
            if token.startswith("**") and token.endswith("**"):
                run = paragraph.add_run(token[2:-2])
                run.font.bold = True
            elif token.startswith("*") and token.endswith("*"):
                run = paragraph.add_run(token[1:-1])
                run.font.italic = True
            elif token.startswith("`") and token.endswith("`"):
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(199, 37, 78)  # Crimson highlight
            elif token.startswith("[") and ")" in token:
                link_match = re.match(r"\[(.*?)\]\((.*?)\)", token)
                if link_match:
                    self._add_hyperlink(
                        paragraph, link_match.group(1), link_match.group(2)
                    )
                else:
                    paragraph.add_run(token)
            else:
                run = paragraph.add_run(token)

            if is_header:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    def _add_hyperlink(self, paragraph, text, url):
        part = paragraph.part
        r_id = part.relate_to(
            url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        )
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        new_run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        c = OxmlElement("w:color")
        c.set(qn("w:val"), "0000FF")
        rPr.append(c)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
        new_run.append(rPr)
        text_node = OxmlElement("w:t")
        text_node.text = text
        new_run.append(text_node)
        hyperlink.append(new_run)
        paragraph._p.append(hyperlink)

    def _set_cell_background(self, cell, color_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color_hex)
        tcPr.append(shd)

    def save(self):
        self.doc.save(self.output_path)
        print(f"Document compiled successfully to: {self.output_path}")


def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    repo_root = os.path.abspath(os.path.join(script_dir, "..", ".."))
    doc_dir = os.path.join(repo_root, "docs")

    files_to_compile = [
        os.path.join(doc_dir, "SYSTEM_ARCHITECTURE.md"),
        os.path.join(doc_dir, "BACKEND_API_REFERENCE.md"),
        os.path.join(doc_dir, "DATABASE_SCHEMA.md"),
        os.path.join(doc_dir, "AUTHENTICATION_AND_AUTHORIZATION.md"),
        os.path.join(doc_dir, "FRONTEND_ARCHITECTURE.md"),
        os.path.join(doc_dir, "DEVELOPER_SETUP_GUIDE.md"),
        os.path.join(doc_dir, "RBAC_Management_Permissions.md"),
        os.path.join(doc_dir, "SCAN_ENGINE_DEEP_DIVE.md"),
        os.path.join(doc_dir, "QA_REGRESSION_STRATEGY.md"),
    ]

    output_docx = os.path.join(doc_dir, "CCure_Audit_PRO_Documentation.docx")
    print(f"Starting compilation to {output_docx}...")

    renderer = DocxRenderer(output_docx)
    renderer.render_cover_page()

    for file_path in files_to_compile:
        if not os.path.isfile(file_path):
            print(f"Skipping missing file: {file_path}")
            continue

        filename = os.path.basename(file_path)
        print(f"Compiling section: {filename}")

        section_title = filename.replace(".md", "").replace("_", " ").title()
        renderer.add_section_header(section_title)

        parser = MarkdownParser(file_path)
        tokens = parser.parse()
        renderer.render_tokens(tokens)
        renderer.doc.add_page_break()

    renderer.save()


if __name__ == "__main__":
    main()
