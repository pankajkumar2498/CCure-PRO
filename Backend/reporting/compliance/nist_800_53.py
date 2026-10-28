"""
reporting/compliance/nist_800_53.py
─────────────────────────────────────────────────────────────────────────────
The NIST SP 800-53 Rev. 5 report, as PDF and DOCX.

Moved unchanged from ``routers/compliance.py`` (G6, C8 of
docs/architecture/MODULE_INDEPENDENCE_ANALYSIS.md): a renderer takes the
controls and their rolled-up verdicts the router read, and returns the
document; it reads no table and imports no compliance code.
"""

import io
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

try:
    import docx
    from docx import Document
except ImportError:
    docx = None
    Document = None


def _generate_nist_800_53_docx(
    controls, findings_map, title, subtitle, scope_label=None
):
    try:
        from docx import Document
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        return io.BytesIO(b"python-docx is not installed.")

    doc = Document()
    p_title = doc.add_paragraph()
    run_t = p_title.add_run(title)
    run_t.font.size = Pt(18)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(10, 46, 92)  # Federal Navy Blue

    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run(subtitle)
    run_sub.font.italic = True
    if scope_label:
        doc.add_paragraph(f"Scope / Baseline: {scope_label}")

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_names = [
        "Control Number",
        "Title",
        "Family",
        "FedRAMP Baseline",
        "Zero Trust",
        "Verdict",
    ]
    for i, name in enumerate(hdr_names):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0A2E5C"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)
        hdr_cells[i].text = name
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    for c in controls:
        f = findings_map.get(c.id)
        verdict = f.status.upper() if f else "UNTESTED"
        baselines = ", ".join(c.saq_profiles) if c.saq_profiles else "N/A"
        zta = "Yes" if "zero_trust" in (c.tags or []) else "No"

        row_cells = table.add_row().cells
        row_cells[0].text = c.control_number
        row_cells[1].text = c.title
        row_cells[2].text = c.theme or "N/A"
        row_cells[3].text = baselines
        row_cells[4].text = zta
        row_cells[5].text = verdict

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _generate_nist_800_53_pdf(
    controls, findings_map, title, subtitle, scope_label=None
):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        "NISTTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0A2E5C"),
    )
    sub_style = ParagraphStyle(
        "NISTSub",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#4b5563"),
    )
    hdr_style = ParagraphStyle(
        "NISTHeader",
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=10,
        textColor=colors.white,
    )
    badge_style = ParagraphStyle(
        "BadgeText",
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9,
        alignment=1,
        textColor=colors.white,
    )
    kpi_style = ParagraphStyle(
        "KPIStyle",
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        alignment=1,
        textColor=colors.HexColor("#334155"),
    )

    story.append(Paragraph(title, title_style))
    story.append(Paragraph(subtitle, sub_style))
    if scope_label:
        story.append(
            Paragraph(f"<b>Filter Baseline:</b> {scope_label}", styles["Normal"])
        )
    story.append(Spacer(1, 10))

    # Executive Headline Summary Box
    total_c = len(controls)
    passed_c = sum(
        1
        for c in controls
        if findings_map.get(c.id) and findings_map.get(c.id).status == "pass"
    )
    failed_c = sum(
        1
        for c in controls
        if findings_map.get(c.id) and findings_map.get(c.id).status == "fail"
    )
    untested_c = total_c - (passed_c + failed_c)
    score_pct = round((passed_c / total_c * 100), 1) if total_c > 0 else 0.0

    kpi_table = Table(
        [
            [
                Paragraph(
                    f"TOTAL CONTROLS<br/><font size=13 color='#0F172A'><b>{total_c}</b></font>",
                    kpi_style,
                ),
                Paragraph(
                    f"PASSED<br/><font size=13 color='#16A34A'><b>{passed_c}</b></font>",
                    kpi_style,
                ),
                Paragraph(
                    f"FAILED<br/><font size=13 color='#DC2626'><b>{failed_c}</b></font>",
                    kpi_style,
                ),
                Paragraph(
                    f"UNTESTED<br/><font size=13 color='#475569'><b>{untested_c}</b></font>",
                    kpi_style,
                ),
                Paragraph(
                    f"POSTURE SCORE<br/><font size=13 color='#0A2E5C'><b>{score_pct}%</b></font>",
                    kpi_style,
                ),
            ]
        ],
        colWidths=[100, 100, 100, 100, 110],
    )
    kpi_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.append(kpi_table)
    story.append(Spacer(1, 12))

    table_data = [
        [
            Paragraph("Control ID", hdr_style),
            Paragraph("Title", hdr_style),
            Paragraph("Family", hdr_style),
            Paragraph("FedRAMP Baseline", hdr_style),
            Paragraph("Zero Trust", hdr_style),
            Paragraph("Verdict", hdr_style),
        ]
    ]
    for c in controls:
        f = findings_map.get(c.id)
        verdict = f.status.upper() if f else "UNTESTED"
        verdict_color = (
            colors.HexColor("#16a34a")
            if verdict == "PASS"
            else (
                colors.HexColor("#dc2626")
                if verdict == "FAIL"
                else colors.HexColor("#6b7280")
            )
        )

        b_str = ", ".join(c.saq_profiles) if c.saq_profiles else "LOW"
        b_bg = (
            "#EF4444"
            if "HIGH" in b_str
            else "#F59E0B" if "MODERATE" in b_str else "#22C55E"
        )
        b_cell = Table(
            [[Paragraph(b_str, badge_style)]], colWidths=[70], rowHeights=[16]
        )
        b_cell.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(b_bg)),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )

        v_cell = Table(
            [[Paragraph(verdict, badge_style)]], colWidths=[72], rowHeights=[16]
        )
        v_cell.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), verdict_color),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )

        zta = "Yes" if "zero_trust" in (c.tags or []) else "No"

        table_data.append(
            [
                Paragraph(f"<b>{c.control_number}</b>", styles["Normal"]),
                Paragraph(c.title, styles["Normal"]),
                Paragraph(c.theme or "N/A", styles["Normal"]),
                b_cell,
                Paragraph(zta, styles["Normal"]),
                v_cell,
            ]
        )

    t = Table(table_data, colWidths=[65, 165, 75, 75, 55, 75])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A2E5C")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, 0), "LEFT"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                ("TOPPADDING", (0, 0), (-1, 0), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e5e7eb")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )

    story.append(t)
    doc.build(story)
    buffer.seek(0)
    return buffer
