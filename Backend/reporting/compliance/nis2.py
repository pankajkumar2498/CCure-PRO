"""
reporting/compliance/nis2.py
─────────────────────────────────────────────────────────────────────────────
The NIS 2 report, as PDF and DOCX.

Moved unchanged from ``routers/compliance.py`` (G6, C8 of
docs/architecture/MODULE_INDEPENDENCE_ANALYSIS.md): a renderer takes the
controls and their rolled-up verdicts the router read, and returns the
document; it reads no table and imports no compliance code.
"""

try:
    import docx
    from docx import Document
except ImportError:
    docx = None
    Document = None


def _generate_nis2_pdf(
    controls,
    findings_map,
    title="EU NIS 2 Directive Compliance Audit Report",
    subtitle="European Union Directive (EU) 2022/2555 Cybersecurity Risk Management Audit",
    scope_label="All Essential & Important Entities",
    entity_type_label="Essential Entity",
    member_state_label="All EU (ENISA Coordinated)",
):
    from io import BytesIO

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError:
        return BytesIO(b"ReportLab is not installed.")

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "NIS2Title",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1E3A8A"),
    )
    sub_style = ParagraphStyle(
        "NIS2Sub",
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#4B5563"),
    )
    section_style = ParagraphStyle(
        "NIS2Section",
        parent=styles["Heading2"],
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#1E3A8A"),
    )
    badge_style = ParagraphStyle(
        "NIS2Badge",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1,
    )

    total_c = len(controls)
    passes = sum(
        1
        for c in controls
        if findings_map.get(c.id) and findings_map.get(c.id).status.lower() == "pass"
    )
    fails = sum(
        1
        for c in controls
        if findings_map.get(c.id) and findings_map.get(c.id).status.lower() == "fail"
    )
    untested = total_c - (passes + fails)
    score_pct = round((passes / total_c * 100), 1) if total_c > 0 else 0.0
    fine_ceiling = (
        "EUR 10,000,000 or 2% Turnover"
        if "Essential" in entity_type_label
        else "EUR 7,000,000 or 1.4% Turnover"
    )

    story = [
        Paragraph(f"<b>{title}</b>", title_style),
        Spacer(1, 4),
        Paragraph(subtitle, sub_style),
        Spacer(1, 4),
        Paragraph(
            f"<b>Entity Scope:</b> {entity_type_label} | <b>Competent Authority:</b> {member_state_label} | <b>Sector Scope:</b> {scope_label}",
            sub_style,
        ),
        Paragraph(
            f"<b>Statutory Penalty Ceiling:</b> {fine_ceiling} | <b>Compliance Rate:</b> {score_pct}%",
            sub_style,
        ),
        Spacer(1, 10),
        Paragraph("<b>1. Executive Summary Statistics</b>", section_style),
        Spacer(1, 4),
    ]

    exec_stats_data = [
        ["Total Assessed Controls", "Pass", "Fail", "Untested", "Compliance Score"],
        [str(total_c), str(passes), str(fails), str(untested), f"{score_pct}%"],
    ]
    t_stats = Table(exec_stats_data, colWidths=[120, 80, 80, 80, 180])
    t_stats.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ]
        )
    )
    story.append(t_stats)
    story.append(Spacer(1, 10))

    story.append(
        Paragraph(
            "<b>2. Article 23 CSIRT Mandatory Incident SLA Timeline</b>", section_style
        )
    )
    story.append(Spacer(1, 4))
    sla_data = [
        ["Phase", "Mandated SLA Timeline", "Scope & Description"],
        [
            "Early Warning",
            "≤ 24 Hours",
            "Notify CSIRT/Competent Authority of significant incident or cyber threat",
        ],
        [
            "Incident Notification",
            "≤ 72 Hours",
            "Submit initial assessment, severity indicator, and IoC indicators",
        ],
        [
            "Final Report",
            "≤ 1 Month",
            "Provide complete root-cause analysis, mitigation, and cross-border impact",
        ],
    ]
    t_sla = Table(sla_data, colWidths=[120, 130, 290])
    t_sla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ]
        )
    )
    story.append(t_sla)
    story.append(Spacer(1, 10))

    story.append(
        Paragraph(
            "<b>3. Article 20 Management Board Governance Attestation</b>",
            section_style,
        )
    )
    story.append(Spacer(1, 4))
    story.append(
        Paragraph(
            '<i>"Under Article 20 of EU NIS 2 Directive (EU 2022/2555), members of the executive management body '
            "have personally reviewed the cybersecurity risk management posture, completed mandatory training, "
            'and accept personal liability for compliance oversight."</i>',
            sub_style,
        )
    )
    story.append(Spacer(1, 10))

    story.append(
        Paragraph(
            "<b>4. Article 21 Detailed Control Evaluation Matrix</b>", section_style
        )
    )
    story.append(Spacer(1, 4))

    table_data = [
        [
            "Control Number",
            "Control Title",
            "Article 21 Domain",
            "Scannability",
            "Verdict",
        ]
    ]

    for c in controls:
        f = findings_map.get(c.id)
        verdict = f.status.upper() if f else "UNTESTED"
        verdict_color = (
            colors.HexColor("#16A34A")
            if verdict == "PASS"
            else (
                colors.HexColor("#DC2626")
                if verdict == "FAIL"
                else colors.HexColor("#6B7280")
            )
        )

        v_cell = Table(
            [[Paragraph(verdict, badge_style)]], colWidths=[65], rowHeights=[14]
        )
        v_cell.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), verdict_color),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )

        table_data.append(
            [
                Paragraph(f"<b>{c.control_number}</b>", styles["Normal"]),
                Paragraph(c.title, styles["Normal"]),
                Paragraph(c.theme or "N/A", styles["Normal"]),
                Paragraph(
                    c.scannability.upper() if c.scannability else "AUTO",
                    styles["Normal"],
                ),
                v_cell,
            ]
        )

    t = Table(table_data, colWidths=[90, 210, 105, 70, 65])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, 0), "LEFT"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                ("TOPPADDING", (0, 0), (-1, 0), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )

    story.append(t)
    doc.build(story)
    buffer.seek(0)
    return buffer


def _generate_nis2_docx(
    controls,
    findings_map,
    title="EU NIS 2 Directive Compliance Audit Report",
    subtitle="European Union Directive (EU) 2022/2555 Cybersecurity Risk Management Audit",
    scope_label="All Essential & Important Entities",
    entity_type_label="Essential Entity",
    member_state_label="All EU (ENISA Coordinated)",
):
    from io import BytesIO

    try:
        from docx import Document
    except ImportError:
        return BytesIO(b"python-docx is not installed.")

    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(subtitle)
    doc.add_paragraph(
        f"Entity Classification: {entity_type_label} | Competent Authority: {member_state_label} | Sector Scope: {scope_label}"
    )

    doc.add_heading("1. Article 23 CSIRT SLA Timeline", level=1)
    doc.add_paragraph(
        "Early Warning: ≤ 24h | Incident Notification: ≤ 72h | Final Report: ≤ 1 Month"
    )

    doc.add_heading("2. Article 20 Board Attestation", level=1)
    doc.add_paragraph(
        "Executive management bodies accept personal accountability and confirm mandatory cybersecurity training under Article 20."
    )

    doc.add_heading("3. Control Evaluation Matrix", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_titles = [
        "Control ID",
        "Control Title",
        "Article 21 Domain",
        "Scannability",
        "Verdict",
    ]
    for i, t in enumerate(hdr_titles):
        hdr_cells[i].text = t

    for c in controls:
        f = findings_map.get(c.id)
        verdict = f.status.upper() if f else "UNTESTED"

        row_cells = table.add_row().cells
        row_cells[0].text = c.control_number
        row_cells[1].text = c.title
        row_cells[2].text = c.theme or "N/A"
        row_cells[3].text = c.scannability.upper() if c.scannability else "AUTO"
        row_cells[4].text = verdict

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
