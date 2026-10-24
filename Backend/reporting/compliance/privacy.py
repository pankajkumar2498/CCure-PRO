"""
reporting/compliance/privacy.py
─────────────────────────────────────────────────────────────────────────────
The privacy reports (DPDP 2023, GDPR), as PDF and DOCX.

Moved unchanged from ``routers/compliance.py`` (G6, C8 of
docs/architecture/MODULE_INDEPENDENCE_ANALYSIS.md): a renderer takes the
controls and their rolled-up verdicts the router read, and returns the
document; it reads no table and imports no compliance code.
"""

import io
from reporting.compliance.common import _verdict_of
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

try:
    import docx
    from docx import Document
except ImportError:
    docx = None
    Document = None


def _generate_privacy_docx(
    controls: list,
    findings_map: dict,
    title: str,
    subtitle: str,
    scope_label: str = "",
) -> io.BytesIO:
    """Generates Word report for privacy frameworks (DPDP / GDPR)."""
    try:
        from docx import Document
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        return io.BytesIO(b"python-docx is not installed.")

    doc = Document()
    p_title = doc.add_paragraph()
    run_t = p_title.add_run(title)
    run_t.font.size = Pt(18)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(29, 78, 216)

    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run(subtitle)
    run_sub.font.italic = True
    if scope_label:
        doc.add_paragraph(f"Scope Filter: {scope_label}")

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_names = [
        "Control Number",
        "Title",
        "Theme / Principle",
        "Scannability",
        "Verdict",
    ]
    for i, name in enumerate(hdr_names):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1D4ED8"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)
        hdr_cells[i].text = name
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    for c in controls:
        row_cells = table.add_row().cells
        row_cells[0].text = str(c.control_number or c.id)
        row_cells[1].text = str(c.title)
        row_cells[2].text = str(c.theme)
        row_cells[3].text = str(c.scannability).upper()
        row_cells[4].text = _verdict_of(findings_map, c.id)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _generate_privacy_pdf(
    controls: list,
    findings_map: dict,
    title: str,
    subtitle: str,
    primary_color: str = "#003366",
    scope_label: str = "",
) -> io.BytesIO:
    """Generates PDF document for privacy frameworks with custom brand color palette."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Title"],
        fontSize=16,
        textColor=colors.HexColor(primary_color),
    )
    cell_style = ParagraphStyle("CellStyle", parent=styles["Normal"], fontSize=8)

    story = [
        Paragraph(title, title_style),
        Paragraph(subtitle, styles["Normal"]),
    ]
    if scope_label:
        story.append(Paragraph(f"<b>Scope Filter:</b> {scope_label}", styles["Normal"]))
    story.append(Spacer(1, 10 * mm))

    table_data = [
        [
            Paragraph("<b>Control ID</b>", cell_style),
            Paragraph("<b>Title</b>", cell_style),
            Paragraph("<b>Theme</b>", cell_style),
            Paragraph("<b>Scannability</b>", cell_style),
            Paragraph("<b>Verdict</b>", cell_style),
        ]
    ]

    for c in controls:
        status_text = _verdict_of(findings_map, c.id)
        table_data.append(
            [
                Paragraph(str(c.control_number or c.id), cell_style),
                Paragraph(str(c.title), cell_style),
                Paragraph(str(c.theme), cell_style),
                Paragraph(str(c.scannability).upper(), cell_style),
                Paragraph(f"<b>{status_text}</b>", cell_style),
            ]
        )

    t = Table(table_data, colWidths=[25 * mm, 75 * mm, 30 * mm, 25 * mm, 25 * mm])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(primary_color)),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(t)

    doc.build(story)
    buffer.seek(0)
    return buffer
