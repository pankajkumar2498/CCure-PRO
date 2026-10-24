"""
reporting/compliance/nist_800_82.py
─────────────────────────────────────────────────────────────────────────────
The NIST SP 800-82 report, as PDF and DOCX.

Moved unchanged from ``routers/compliance.py`` (G6, C8 of
docs/architecture/MODULE_INDEPENDENCE_ANALYSIS.md): a renderer takes the
controls and their rolled-up verdicts the router read, and returns the
document; it reads no table and imports no compliance code.
"""

try:
    import docx
    from docx import Document
except ImportError:
    docx = None
    Document = None


def _generate_nist_800_82_pdf(
    controls,
    findings_map,
    title="NIST SP 800-82 Rev. 3 OT Security Assessment Report",
    subtitle="Operational Technology (OT) Security Control Ledger & Purdue Alignment",
    scope_label="All OT Assets & Purdue Levels (L0–L5)",
):
    from io import BytesIO

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError:
        return BytesIO(b"ReportLab is not installed.")

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    story = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "NIST82Title",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#2563EB"),
        alignment=0,
    )
    sub_style = ParagraphStyle(
        "NIST82Sub",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#6B7280"),
    )
    badge_style = ParagraphStyle(
        "NIST82Badge",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1,
    )

    story.append(Paragraph(title, title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(subtitle, sub_style))
    if scope_label:
        story.append(Spacer(1, 4))
        story.append(
            Paragraph(f"<b>Scope Baseline:</b> {scope_label}", styles["Normal"])
        )
    story.append(Spacer(1, 14))

    table_data = [
        [
            "Control ID",
            "Control Title",
            "Theme",
            "Purdue Level",
            "Components",
            "Verdict",
        ]
    ]
    purdue_levels_set = {"L0", "L1", "L2", "L3", "L3.5", "L4", "L5"}
    for c in controls:
        f = findings_map.get(c.id)
        verdict = f.status.upper() if f else "UNTESTED"
        verdict_color = (
            colors.HexColor("#16A34A")
            if verdict == "PASS"
            else (
                colors.HexColor("#DC2626")
                if verdict == "FAIL"
                else colors.HexColor("#6B7280")
            )
        )

        tags = c.tags or []
        purdue_str = ", ".join([t for t in tags if t in purdue_levels_set]) or "L0-L5"
        comp_types = {
            "PLC",
            "RTU",
            "SCADA",
            "DCS",
            "SIS",
            "HMI",
            "HISTORIAN",
            "IED",
            "NDD",
            "EDR",
            "HDR",
            "SW",
        }
        comp_str = ", ".join([t for t in tags if t in comp_types]) or "N/A"

        v_cell = Table(
            [[Paragraph(verdict, badge_style)]], colWidths=[55], rowHeights=[14]
        )
        v_cell.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), verdict_color),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )

        table_data.append(
            [
                Paragraph(f"<b>{c.control_number}</b>", styles["Normal"]),
                Paragraph(c.title, styles["Normal"]),
                Paragraph(c.theme or "N/A", styles["Normal"]),
                Paragraph(purdue_str, styles["Normal"]),
                Paragraph(comp_str, styles["Normal"]),
                v_cell,
            ]
        )

    t = Table(table_data, colWidths=[80, 180, 75, 75, 70, 60])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, 0), "LEFT"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                ("TOPPADDING", (0, 0), (-1, 0), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )

    story.append(t)
    doc.build(story)
    buffer.seek(0)
    return buffer


def _generate_nist_800_82_docx(
    controls,
    findings_map,
    title="NIST SP 800-82 Rev. 3 OT Security Assessment Report",
    subtitle="Operational Technology (OT) Security Control Ledger & Purdue Alignment",
    scope_label="All OT Assets & Purdue Levels (L0–L5)",
):
    from io import BytesIO

    try:
        from docx import Document
    except ImportError:
        return BytesIO(b"python-docx is not installed.")

    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(subtitle)
    if scope_label:
        doc.add_paragraph(f"Scope Baseline: {scope_label}")

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_titles = [
        "Control ID",
        "Control Title",
        "Theme",
        "Purdue Level",
        "Components",
        "Verdict",
    ]
    for i, t in enumerate(hdr_titles):
        hdr_cells[i].text = t

    purdue_levels_set = {"L0", "L1", "L2", "L3", "L3.5", "L4", "L5"}
    comp_types = {
        "PLC",
        "RTU",
        "SCADA",
        "DCS",
        "SIS",
        "HMI",
        "HISTORIAN",
        "IED",
        "NDD",
        "EDR",
        "HDR",
        "SW",
    }
    for c in controls:
        f = findings_map.get(c.id)
        verdict = f.status.upper() if f else "UNTESTED"
        tags = c.tags or []
        purdue_str = ", ".join([t for t in tags if t in purdue_levels_set]) or "L0-L5"
        comp_str = ", ".join([t for t in tags if t in comp_types]) or "N/A"

        row_cells = table.add_row().cells
        row_cells[0].text = c.control_number
        row_cells[1].text = c.title
        row_cells[2].text = c.theme or "N/A"
        row_cells[3].text = purdue_str
        row_cells[4].text = comp_str
        row_cells[5].text = verdict

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
