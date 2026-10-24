"""
reporting/compliance/pci_dss.py
─────────────────────────────────────────────────────────────────────────────
The PCI DSS v4.0.1 report, as PDF and DOCX.

Moved unchanged from ``routers/compliance.py`` (G6, C8 of
docs/architecture/MODULE_INDEPENDENCE_ANALYSIS.md): a renderer takes the
controls and their rolled-up verdicts the router read, and returns the
document; it reads no table and imports no compliance code.
"""

import io
from reporting.compliance.common import _verdict_of
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

try:
    import docx
    from docx import Document
except ImportError:
    docx = None
    Document = None


def _generate_pci_dss_docx(
    controls: list,
    findings_map: dict,
    title: str = "PCI-DSS v4.0 Compliance Assessment Report",
    subtitle: str = "Payment Card Industry Data Security Standard Assessment Ledger",
) -> io.BytesIO:
    """Generates the Word PCI-DSS report using python-docx."""
    try:
        from docx import Document
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        return io.BytesIO(b"python-docx is not installed.")

    doc = Document()
    p_title = doc.add_paragraph()
    run_t = p_title.add_run(title)
    run_t.font.size = Pt(18)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(30, 64, 175)

    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run(subtitle)
    run_sub.font.italic = True

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_names = [
        "Control ID",
        "Name",
        "SAQ Profile Relevance",
        "March 31, 2025 Mandatory",
        "CDE Scope Tagging",
        "Verdict Status",
    ]
    for i, name in enumerate(hdr_names):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E40AF"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)
        hdr_cells[i].text = name
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    for c in controls:
        row_cells = table.add_row().cells
        row_cells[0].text = str(c.control_number or c.id)
        row_cells[1].text = str(c.title)

        saq_str = ", ".join(c.saq_profiles) if c.saq_profiles else "All"
        row_cells[2].text = saq_str

        is_mand = (
            "Yes"
            if (
                c.tags
                and "v4_2025_mandatory" in c.tags
                or getattr(c, "v4_2025_mandatory", False)
            )
            else "No"
        )
        row_cells[3].text = is_mand

        is_cde = "Yes" if (c.tags and "cde" in [t.lower() for t in c.tags]) else "No"
        row_cells[4].text = is_cde

        f = findings_map.get(c.id)
        row_cells[5].text = str(f.status).upper() if f else "PENDING"

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _generate_pci_dss_pdf(
    controls: list,
    findings_map: dict,
    title: str = "PCI-DSS v4.0 Compliance Assessment Report",
    subtitle: str = "Payment Card Industry Data Security Standard Assessment Ledger",
) -> io.BytesIO:
    """Generates the PCI-DSS PDF document using ReportLab."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Title"],
        fontSize=16,
        textColor=colors.HexColor("#003366"),
    )
    cell_style = ParagraphStyle("CellStyle", parent=styles["Normal"], fontSize=8)

    story = [
        Paragraph(title, title_style),
        Paragraph(subtitle, styles["Normal"]),
        Spacer(1, 10 * mm),
    ]

    table_data = [
        [
            Paragraph("<b>Control ID</b>", cell_style),
            Paragraph("<b>Name</b>", cell_style),
            Paragraph("<b>SAQ Profile Relevance</b>", cell_style),
            Paragraph("<b>March 31, 2025 Mandatory</b>", cell_style),
            Paragraph("<b>CDE Scope Tagging</b>", cell_style),
            Paragraph("<b>Verdict</b>", cell_style),
        ]
    ]

    for c in controls:
        status_text = _verdict_of(findings_map, c.id)

        saq_str = ", ".join(c.saq_profiles) if c.saq_profiles else "All"
        is_mand = (
            "Yes"
            if (
                c.tags
                and "v4_2025_mandatory" in c.tags
                or getattr(c, "v4_2025_mandatory", False)
            )
            else "No"
        )
        is_cde = "Yes" if (c.tags and "cde" in [t.lower() for t in c.tags]) else "No"

        table_data.append(
            [
                Paragraph(c.control_number or c.id, cell_style),
                Paragraph(
                    c.title[:35] + "..." if len(c.title) > 35 else c.title, cell_style
                ),
                Paragraph(saq_str, cell_style),
                Paragraph(is_mand, cell_style),
                Paragraph(is_cde, cell_style),
                Paragraph(
                    (
                        f"<font color='green'>{status_text}</font>"
                        if status_text == "PASS"
                        else f"<font color='red'>{status_text}</font>"
                    ),
                    cell_style,
                ),
            ]
        )

    tbl = Table(
        table_data, colWidths=[15 * mm, 50 * mm, 35 * mm, 25 * mm, 25 * mm, 30 * mm]
    )
    tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#003366")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(tbl)

    doc.build(story)
    buffer.seek(0)
    return buffer
