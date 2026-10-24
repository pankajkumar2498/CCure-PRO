"""
reporting/compliance/iec_62443.py
─────────────────────────────────────────────────────────────────────────────
The IEC 62443 report, as PDF and DOCX.

Moved unchanged from ``routers/compliance.py`` (G6, C8 of
docs/architecture/MODULE_INDEPENDENCE_ANALYSIS.md): a renderer takes the
controls and their rolled-up verdicts the router read, and returns the
document; it reads no table and imports no compliance code.
"""

try:
    import docx
    from docx import Document
except ImportError:
    docx = None
    Document = None


# ─────────────────────────────────────────────────────────────────────────────
# IEC 62443 Executive Audit Report Generators
# ─────────────────────────────────────────────────────────────────────────────


def _generate_iec_62443_pdf(
    controls,
    findings_map,
    title="IEC 62443 Industrial Cybersecurity Audit Report",
    subtitle="IACS System Security & Security Level Readiness Ledger",
    scope_label="All Security Levels (SL1–SL4)",
):
    from io import BytesIO

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        "IECTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#F43F5E"),
        alignment=0,
    )
    sub_style = ParagraphStyle(
        "IECSub",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#6B7280"),
    )
    hdr_style = ParagraphStyle(
        "IECHdr",
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=10,
        textColor=colors.white,
    )
    badge_style = ParagraphStyle(
        "IECBadge",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9,
        textColor=colors.white,
        alignment=1,
    )
    kpi_style = ParagraphStyle(
        "IECKPIStyle",
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        alignment=1,
        textColor=colors.HexColor("#334155"),
    )

    story.append(Paragraph(title, title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(subtitle, sub_style))
    if scope_label:
        story.append(Spacer(1, 4))
        story.append(
            Paragraph(f"<b>Scope Baseline:</b> {scope_label}", styles["Normal"])
        )
    story.append(Spacer(1, 10))

    # Executive Headline Summary Box
    total_c = len(controls)
    passed_c = sum(
        1
        for c in controls
        if findings_map.get(c.id) and findings_map.get(c.id).status == "pass"
    )
    failed_c = sum(
        1
        for c in controls
        if findings_map.get(c.id) and findings_map.get(c.id).status == "fail"
    )
    untested_c = total_c - (passed_c + failed_c)
    score_pct = round((passed_c / total_c * 100), 1) if total_c > 0 else 0.0

    kpi_table = Table(
        [
            [
                Paragraph(
                    f"TOTAL CONTROLS<br/><font size=13 color='#0F172A'><b>{total_c}</b></font>",
                    kpi_style,
                ),
                Paragraph(
                    f"PASSED<br/><font size=13 color='#16A34A'><b>{passed_c}</b></font>",
                    kpi_style,
                ),
                Paragraph(
                    f"FAILED<br/><font size=13 color='#DC2626'><b>{failed_c}</b></font>",
                    kpi_style,
                ),
                Paragraph(
                    f"UNTESTED<br/><font size=13 color='#475569'><b>{untested_c}</b></font>",
                    kpi_style,
                ),
                Paragraph(
                    f"READY SCORE<br/><font size=13 color='#F43F5E'><b>{score_pct}%</b></font>",
                    kpi_style,
                ),
            ]
        ],
        colWidths=[100, 100, 100, 100, 140],
    )
    kpi_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFF1F2")),
                ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#FECDD3")),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#FFE4E6")),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.append(kpi_table)
    story.append(Spacer(1, 12))

    table_data = [
        [
            Paragraph("Requirement ID", hdr_style),
            Paragraph("Control Title", hdr_style),
            Paragraph("Theme", hdr_style),
            Paragraph("Purdue Level", hdr_style),
            Paragraph("Target SL", hdr_style),
            Paragraph("Verdict", hdr_style),
        ]
    ]
    for c in controls:
        f = findings_map.get(c.id)
        verdict = f.status.upper() if f else "UNTESTED"
        verdict_color = (
            colors.HexColor("#16A34A")
            if verdict == "PASS"
            else (
                colors.HexColor("#DC2626")
                if verdict == "FAIL"
                else colors.HexColor("#6B7280")
            )
        )

        tags = c.tags or []
        sl_str = next((t for t in tags if t.startswith("SL")), "SL1-SL4")
        purdue_levels_set = {"L0", "L1", "L2", "L3", "L3.5", "L4", "L5"}
        purdue_str = ", ".join([t for t in tags if t in purdue_levels_set]) or "L1-L3.5"

        v_cell = Table(
            [[Paragraph(verdict, badge_style)]], colWidths=[72], rowHeights=[16]
        )
        v_cell.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), verdict_color),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )

        table_data.append(
            [
                Paragraph(f"<b>{c.control_number}</b>", styles["Normal"]),
                Paragraph(c.title, styles["Normal"]),
                Paragraph(c.theme or "N/A", styles["Normal"]),
                Paragraph(purdue_str, styles["Normal"]),
                Paragraph(sl_str, styles["Normal"]),
                v_cell,
            ]
        )

    t = Table(table_data, colWidths=[80, 175, 70, 75, 60, 75])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F43F5E")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, 0), "LEFT"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                ("TOPPADDING", (0, 0), (-1, 0), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )

    story.append(t)
    doc.build(story)
    buffer.seek(0)
    return buffer


def _generate_iec_62443_docx(
    controls,
    findings_map,
    title="IEC 62443 Industrial Cybersecurity Audit Report",
    subtitle="IACS System Security & Security Level Readiness Ledger",
    scope_label="All Security Levels (SL1–SL4)",
):
    from io import BytesIO

    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        return BytesIO(b"python-docx is not installed.")

    doc = Document()
    doc.add_heading(title, level=0)
    p_sub = doc.add_paragraph(subtitle)
    if scope_label:
        p_scope = doc.add_paragraph(f"Scope Baseline: {scope_label}")

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_titles = [
        "Requirement ID",
        "Control Title",
        "Theme",
        "Purdue Level",
        "Target SL",
        "Verdict",
    ]
    for i, t in enumerate(hdr_titles):
        hdr_cells[i].text = t

    for c in controls:
        f = findings_map.get(c.id)
        verdict = f.status.upper() if f else "UNTESTED"
        tags = c.tags or []
        sl_str = next((t for t in tags if t.startswith("SL")), "SL1-SL4")
        purdue_levels_set = {"L0", "L1", "L2", "L3", "L3.5", "L4", "L5"}
        purdue_str = ", ".join([t for t in tags if t in purdue_levels_set]) or "L1-L3.5"

        row_cells = table.add_row().cells
        row_cells[0].text = c.control_number
        row_cells[1].text = c.title
        row_cells[2].text = c.theme or "N/A"
        row_cells[3].text = purdue_str
        row_cells[4].text = sl_str
        row_cells[5].text = verdict

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
